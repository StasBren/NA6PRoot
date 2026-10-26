#!/usr/bin/env python3
"""Build a ONE-SLIDE ROOT-style PI presentation from real 20-vs-500 CSV.

Reads the previously computed, identical-exposure J/psi MS0
phi_comparison.csv; launches a ROOT C++ macro making two vertically
stacked TH1D histograms (20 above, 500 below), matched x/y scales and
individual 1-sigma Garfield-LUT errors. Writes only PNG and PPTX.
The histograms are rendered using ROOT but no .root/.pdf files are saved.

No Garfield, Geant4, Heed or frozen 720-fold jobs are run or changed.
"""
from __future__ import annotations

import argparse
import csv
import math
import shutil
import subprocess
from pathlib import Path

HERE=Path(__file__).resolve().parent
REPO=HERE.parents[1]
INPUT=REPO/"test_runs/mwpc_stage_c/pi_pitch_fieldwire_20_vs_500_expected_v1/phi_comparison.csv"
OUTPUT=REPO/"test_runs/mwpc_stage_c/pi_lut_root_slide_v1"
MACRO=HERE/"pi_lut_root_slide.C"
GEOMS=("p3p0_field75p0um","p3p0_anode_only")


def validate_csv(path,geometry):
    with path.open(newline="",encoding="utf-8") as stream:
        r=csv.DictReader(stream)
        expected={"geometry","phi_lower_deg","phi_upper_deg",
                  "p20","se20","p500","se500"}
        if not expected.issubset(r.fieldnames or []):
            raise ValueError("Missing phi_comparison.csv fields")
        selected=[row for row in r if row["geometry"]==geometry]
    if len(selected)!=36:
        raise ValueError(f"Expected 36 angular bins, got {len(selected)}")
    sums={"p20":0.,"p500":0.}
    for i,row in enumerate(sorted(selected,key=lambda z:float(z["phi_lower_deg"]))):
        lo=float(row["phi_lower_deg"])
        hi=float(row["phi_upper_deg"])
        if not math.isclose(lo,-180.+i*10.,abs_tol=1e-6) or not \
           math.isclose(hi,lo+10.,abs_tol=1e-6):
            raise ValueError(f"Invalid angular bin {i}")
        for key in ("p20","se20","p500","se500"):
            v=float(row[key])
            if not math.isfinite(v) or v<0:
                raise ValueError(f"Invalid {key} in bin {i}")
        sums["p20"]+=float(row["p20"])
        sums["p500"]+=float(row["p500"])
    if any(abs(sums[name]-1.)>1e-4 for name in sums):
        raise ValueError(f"PDFs do not normalize: {sums}")
    return sums


def require_pptx():
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is required for the PPTX output. "
            "On Ubuntu 24.04, use a virtual environment rather than "
            "sudo apt install python3-pptx: "
            "python3 -m venv --system-site-packages "
            "test_runs/mwpc_stage_c/.venv_pi_slides; "
            "test_runs/mwpc_stage_c/.venv_pi_slides/bin/python "
            "-m pip install python-pptx; "
            "then run the script with that venv's Python."
        ) from exc
    return Presentation, Inches


def save_pptx(png,out,presentation,inches):
    ppt=presentation()
    ppt.slide_width=inches(13.333333)
    ppt.slide_height=inches(7.5)
    slide=ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_picture(str(png),0,0,width=ppt.slide_width,
                             height=ppt.slide_height)
    ppt.save(str(out))
    if not out.is_file() or out.stat().st_size==0:
        raise RuntimeError("PowerPoint slide was not saved")


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv",type=Path,default=INPUT,
                        help="Original generated phi_comparison.csv")
    parser.add_argument("--geometry",choices=GEOMS,
                        default="p3p0_field75p0um")
    parser.add_argument("--out",type=Path,default=OUTPUT)
    parser.add_argument("--validate-only",action="store_true",
                        help="Check measured input without requiring ROOT")
    a=parser.parse_args()
    if not a.csv.is_file():
        raise FileNotFoundError(
            f"Missing real angular data: {a.csv}\n"
            "First run pi_lut_20_vs_500_expected.py --station 0 --chunks 2")
    normalizations=validate_csv(a.csv,a.geometry)
    print(f"PASS: measured CSV, 36 normalized angular bins, "
          f"geometry={a.geometry}; {normalizations}",flush=True)
    if a.validate_only:
        return
    if not MACRO.is_file():
        raise FileNotFoundError(f"Missing ROOT plotting macro: {MACRO}")
    presentation,inches=require_pptx()
    root=shutil.which("root")
    if not root:
        raise EnvironmentError("ROOT executable 'root' not on PATH")
    a.out.mkdir(parents=True,exist_ok=True)
    prefix=a.geometry+"_slide"
    png=a.out/(prefix+".png")
    pptx=a.out/(prefix+".pptx")
    # Never use stale successful results to mask rendering failures.
    for target in (png,pptx):
        if target.exists():
            target.unlink()
    # Standard WSL POSIX paths; no arbitrary shell invocation.
    def quote_root(p):
        s=str(p.resolve())
        if any(char in s for char in ('"',"\n","\r")):
            raise ValueError("Unexpected filename characters")
        return '"'+s+'"'
    command=(str(MACRO.resolve())+"("+
             quote_root(a.csv)+","+quote_root(a.out)+
             ',"'+a.geometry+'")')
    process=subprocess.run(
        [root,"-l","-b","-q",command],cwd=REPO,
        stdout=subprocess.PIPE,stderr=subprocess.STDOUT,
        text=True,check=False)
    print(process.stdout,flush=True)
    if process.returncode or "PASS: ROOT-rendered TH1D" not in process.stdout or \
       not png.is_file() or png.stat().st_size==0:
        raise RuntimeError("ROOT PNG rendering failed; see preceding log")
    save_pptx(png,pptx,presentation,inches)
    print("PASS: one PI slide from measured ROOT histograms and LUT errors")
    print("OUTPUT",png,flush=True)
    print("OUTPUT",pptx,flush=True)


if __name__=="__main__":
    main()
